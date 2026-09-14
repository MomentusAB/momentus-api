"""Skriver en exporterad lagerlista till Excel, CSV, PDF eller Word.

Modulen vet ingenting om databasen eller Tk - den får en färdig
ExportSpec (kolumner, rader i rätt ordning, inställningar) från
exportdialogen i ui_export.py och skriver en fil. Den ändrar aldrig något.

Paket som behövs (pip install ...):
    openpyxl      -> Excel (.xlsx)
    reportlab     -> PDF
    python-docx   -> Word (.docx)
CSV använder bara Pythons standardbibliotek.
"""

import csv
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path


# Filtyper i Spara-dialogen. Etiketten används för att lista ut vilket format
# användaren valde om sökvägen kommer tillbaka utan filändelse.
FILETYPES = [
    ("Excel", "*.xlsx"),
    ("CSV", "*.csv"),
    ("PDF", "*.pdf"),
    ("Word", "*.docx"),
]
LABEL_BY_EXT = {".xlsx": "Excel", ".csv": "CSV", ".pdf": "PDF", ".docx": "Word"}
EXT_BY_LABEL = {label: ext for ext, label in LABEL_BY_EXT.items()}

# Kolumner som är tal - högerställs i PDF/Word och skrivs som tal i Excel.
NUMERIC_KEYS = {"quantity", "unit_cost"}

# Rubriker och relativa bredder för de kolumner som inte finns i COLUMN_DEFINITIONS
# (huvud-/underkategori finns bara som egna kolumner i platt läge).
PRODUCT_LABEL = "Produkt"
PRODUCT_WIDTH = 230
MAIN_LABEL, MAIN_WIDTH = "Huvudkategori", 140
SUB_LABEL, SUB_WIDTH = "Underkategori", 130


class MissingDependency(Exception):
    """Ett paket som behövs för formatet är inte installerat."""


@dataclass
class ExportSpec:
    title: str
    subtitle_lines: list                # beskrivning av urvalet, en rad per post
    columns: list                       # [(key, label, width, anchor), ...] i önskad ordning
    rows: list                          # [{"_main": ..., "_sub": ..., "product_name": ..., key: värde}, ...] i önskad ordning
    grouped: bool = True                # rubrikrader per kategori (PDF/Word). Excel/CSV är alltid platta.
    landscape: bool = True              # PDF/Word
    page_break_per_main: bool = False   # PDF/Word
    summary: dict = None                # {"count": n, "total_qty": q, "value_by_currency": {"SEK": 123.0}} eller None
    exported_at: str = field(default_factory=lambda: datetime.now().strftime("%Y-%m-%d %H:%M"))


def resolve_path(path: str, chosen_type_label: str, default_ext: str) -> str:
    """Ser till att sökvägen har en filändelse vi kan hantera.

    Windows lägger själv på ändelsen för vald filtyp, andra plattformar inte
    alltid - då används filtypen användaren valde i dialogen, annars default."""
    ext = Path(path).suffix.lower()
    if ext in LABEL_BY_EXT:
        return path
    ext = EXT_BY_LABEL.get((chosen_type_label or "").strip(), default_ext)
    if ext not in LABEL_BY_EXT:
        ext = ".xlsx"
    return str(path) + ext


def export_to_file(path: str, spec: ExportSpec) -> None:
    ext = Path(path).suffix.lower()
    if ext == ".xlsx":
        _export_xlsx(path, spec)
    elif ext == ".csv":
        _export_csv(path, spec)
    elif ext == ".pdf":
        _export_pdf(path, spec)
    elif ext == ".docx":
        _export_docx(path, spec)
    else:
        raise ValueError(f"Okänt filformat: {ext or '(ingen filändelse)'}")


# ---------------------------------------------------------------- hjälpfunktioner

def format_value(key, value, decimal_comma=False) -> str:
    """Textrepresentation av ett cellvärde för CSV/PDF/Word."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Ja" if value else "Nej"
    if isinstance(value, float):
        text = f"{value:.2f}"
        return text.replace(".", ",") if decimal_comma else text
    if isinstance(value, int):
        return str(value)
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value)


def summary_lines(summary: dict) -> list:
    lines = [
        f"Antal artiklar: {summary.get('count', 0)}",
        f"Summa antal: {summary.get('total_qty', 0)}",
    ]
    values = summary.get("value_by_currency") or {}
    if values:
        parts = [f"{total:,.2f} {cur}".replace(",", " ") for cur, total in values.items()]
        lines.append("Lagervärde: " + "  |  ".join(parts))
    else:
        lines.append("Lagervärde: -")
    return lines


def _flat_headers(spec: ExportSpec):
    return [MAIN_LABEL, SUB_LABEL, PRODUCT_LABEL] + [label for _k, label, *_ in spec.columns]


def _flat_row_values(spec: ExportSpec, row: dict):
    return [row.get("_main", ""), row.get("_sub", ""), row.get("product_name", "")] + [
        row.get(key) for key, *_ in spec.columns
    ]


# ---------------------------------------------------------------- CSV

def _export_csv(path: str, spec: ExportSpec) -> None:
    # Semikolon + UTF-8 med BOM + decimalkomma: öppnas direkt i svenskt Excel
    # med rätt kolumner, å/ä/ö och tal som tal.
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(_flat_headers(spec))
        for row in spec.rows:
            values = _flat_row_values(spec, row)
            writer.writerow([format_value(None, v, decimal_comma=True) for v in values])

        if spec.summary:
            writer.writerow([])
            for line in summary_lines(spec.summary):
                label, _, value = line.partition(": ")
                writer.writerow([label, value])


# ---------------------------------------------------------------- Excel

def _export_xlsx(path: str, spec: ExportSpec) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise MissingDependency(f"Excel-export kräver paketet openpyxl (pip install openpyxl).\n\nImportfel: {e}") from e

    wb = Workbook()
    ws = wb.active
    ws.title = "Lager"

    headers = _flat_headers(spec)
    ws.append(headers)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E6E6E6")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    keys = [None, None, None] + [key for key, *_ in spec.columns]
    widths = [len(h) for h in headers]

    for row in spec.rows:
        values = _flat_row_values(spec, row)
        ws.append(values)
        for i, v in enumerate(values):
            widths[i] = max(widths[i], len(format_value(keys[i], v)))

    last_data_row = ws.max_row
    last_col_letter = get_column_letter(len(headers))

    # Talformat och högerställning per kolumn
    for col_idx, key in enumerate(keys, start=1):
        if key == "unit_cost":
            fmt, align = "#,##0.00", "right"
        elif key == "quantity":
            fmt, align = "0", "right"
        elif key == "last_inventory_check":
            fmt, align = "yyyy-mm-dd", "center"
        else:
            continue
        for (cell,) in ws.iter_rows(min_row=2, max_row=last_data_row, min_col=col_idx, max_col=col_idx):
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal=align)

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(8, min(45, w + 2))

    ws.freeze_panes = "A2"
    if last_data_row >= 2:
        ws.auto_filter.ref = f"A1:{last_col_letter}{last_data_row}"

    if spec.summary:
        ws.append([])
        for line in summary_lines(spec.summary):
            label, _, value = line.partition(": ")
            ws.append([label, value])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    # Ett informationsblad med vad exporten innehåller
    info = wb.create_sheet("Info")
    info.append([spec.title])
    info["A1"].font = Font(bold=True, size=12)
    info.append([f"Exporterad {spec.exported_at}"])
    info.append([])
    for line in spec.subtitle_lines:
        info.append([line])
    info.column_dimensions["A"].width = 90

    wb.save(path)


# ---------------------------------------------------------------- PDF

def _export_pdf(path: str, spec: ExportSpec) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_RIGHT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as e:
        raise MissingDependency(f"PDF-export kräver paketet reportlab (pip install reportlab).\n\nImportfel: {e}") from e
    from xml.sax.saxutils import escape

    pagesize = landscape(A4) if spec.landscape else A4
    doc = SimpleDocTemplate(
        path,
        pagesize=pagesize,
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title=spec.title,
    )

    styles = getSampleStyleSheet()
    # wordWrap="CJK" låter långa strängar utan mellanslag (streckkoder,
    # artikelnummer) brytas inne i cellen istället för att sticka ut.
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontName="Helvetica", fontSize=7, leading=8.5, wordWrap="CJK")
    cell_right = ParagraphStyle("cell_right", parent=cell, alignment=TA_RIGHT)
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold")
    main_style = ParagraphStyle("main", parent=cell, fontName="Helvetica-Bold", fontSize=8.5, leading=10)
    sub_style = ParagraphStyle("sub", parent=cell, fontName="Helvetica-Bold")
    title_style = ParagraphStyle("title", parent=styles["Heading1"], fontSize=14, leading=17, spaceAfter=2)
    meta_style = ParagraphStyle("meta", parent=styles["Normal"], fontSize=8, leading=10, textColor=colors.HexColor("#555555"))

    def P(text, style):
        return Paragraph(escape(text or ""), style)

    flat = not spec.grouped
    col_keys = [key for key, *_ in spec.columns]
    if flat:
        headers = [MAIN_LABEL, SUB_LABEL, PRODUCT_LABEL] + [label for _k, label, *_ in spec.columns]
        weights = [MAIN_WIDTH, SUB_WIDTH, PRODUCT_WIDTH] + [width for _k, _l, width, _a in spec.columns]
    else:
        headers = [PRODUCT_LABEL] + [label for _k, label, *_ in spec.columns]
        weights = [PRODUCT_WIDTH] + [width for _k, _l, width, _a in spec.columns]
    ncols = len(headers)
    col_widths = [doc.width * w / sum(weights) for w in weights]

    base_cmds = [
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#999999")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6e6e6")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]

    def data_cells(row):
        cells = []
        if flat:
            cells.append(P(row.get("_main", ""), cell))
            cells.append(P(row.get("_sub", ""), cell))
        cells.append(P(row.get("product_name", ""), cell))
        for key in col_keys:
            style = cell_right if key in NUMERIC_KEYS else cell
            cells.append(P(format_value(key, row.get(key)), style))
        return cells

    def build_table(block):
        """block: lista av ("main", text) / ("sub", text) / ("row", raddict)."""
        data = [[P(h, head) for h in headers]]
        cmds = list(base_cmds)
        for kind, payload in block:
            r = len(data)
            if kind == "main":
                data.append([P(payload, main_style)] + [""] * (ncols - 1))
                cmds += [("SPAN", (0, r), (-1, r)), ("BACKGROUND", (0, r), (-1, r), colors.HexColor("#d9d9d9"))]
            elif kind == "sub":
                data.append([P(payload, sub_style)] + [""] * (ncols - 1))
                cmds += [("SPAN", (0, r), (-1, r)), ("BACKGROUND", (0, r), (-1, r), colors.HexColor("#f2f2f2"))]
            else:
                data.append(data_cells(payload))
        table = Table(data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle(cmds))
        return table

    story = [P(spec.title, title_style), P(f"Exporterad {spec.exported_at}", meta_style)]
    for line in spec.subtitle_lines:
        story.append(P(line, meta_style))
    story.append(Spacer(1, 3 * mm))

    if flat:
        story.append(build_table([("row", row) for row in spec.rows]))
    else:
        blocks = []
        current_main = current_sub = None
        block = None
        for row in spec.rows:
            if row.get("_main") != current_main or block is None:
                current_main, current_sub = row.get("_main"), None
                block = [("main", current_main or "")]
                blocks.append(block)
            if row.get("_sub") != current_sub:
                current_sub = row.get("_sub")
                block.append(("sub", current_sub or ""))
            block.append(("row", row))
        for i, blk in enumerate(blocks):
            if i > 0:
                story.append(PageBreak() if spec.page_break_per_main else Spacer(1, 4 * mm))
            story.append(build_table(blk))

    if spec.summary:
        story.append(Spacer(1, 4 * mm))
        for line in summary_lines(spec.summary):
            story.append(P(line, meta_style))

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(10 * mm, 7 * mm, f"{spec.title} - exporterad {spec.exported_at}")
        canvas.drawRightString(pagesize[0] - 10 * mm, 7 * mm, f"Sida {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


# ---------------------------------------------------------------- Word

def _export_docx(path: str, spec: ExportSpec) -> None:
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor
        from docx.table import _Cell
    except ImportError as e:
        raise MissingDependency(f"Word-export kräver paketet python-docx (pip install python-docx).\n\nImportfel: {e}") from e

    doc = Document()
    section = doc.sections[0]
    if spec.landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Mm(12)
    section.top_margin = section.bottom_margin = Mm(12)
    usable_mm = (section.page_width - section.left_margin - section.right_margin) / Mm(1)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(8)

    doc.add_heading(spec.title, level=1)
    meta_lines = [f"Exporterad {spec.exported_at}"] + list(spec.subtitle_lines)
    for line in meta_lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        p.paragraph_format.space_after = Pt(0)

    flat = not spec.grouped
    col_keys = [key for key, *_ in spec.columns]
    if flat:
        headers = [MAIN_LABEL, SUB_LABEL, PRODUCT_LABEL] + [label for _k, label, *_ in spec.columns]
        weights = [MAIN_WIDTH, SUB_WIDTH, PRODUCT_WIDTH] + [width for _k, _l, width, _a in spec.columns]
    else:
        headers = [PRODUCT_LABEL] + [label for _k, label, *_ in spec.columns]
        weights = [PRODUCT_WIDTH] + [width for _k, _l, width, _a in spec.columns]
    ncols = len(headers)
    col_widths_mm = [usable_mm * w / sum(weights) for w in weights]
    first_data_col = ncols - len(col_keys)
    numeric_cols = {first_data_col + i for i, key in enumerate(col_keys) if key in NUMERIC_KEYS}

    def shade(cell, fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

    def set_text(cell, text, bold=False, right=False, size=8):
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(text or "")
        run.bold = bold
        run.font.size = Pt(size)
        if right:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

    def row_cells(table, tr):
        # Direkt via XML: table.rows[i].cells bygger om hela cellistan varje
        # gång och blir mycket långsamt för stora tabeller.
        return [_Cell(tc, table) for tc in tr.tc_lst]

    def mark_header_row(tr):
        tr_pr = tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        tr_pr.append(el)

    def new_table():
        table = doc.add_table(rows=1, cols=ncols)
        table.style = "Table Grid"
        table.autofit = False
        header_tr = table._tbl.tr_lst[0]
        for cell, label in zip(row_cells(table, header_tr), headers):
            set_text(cell, label, bold=True)
            shade(cell, "E6E6E6")
        mark_header_row(header_tr)
        return table

    def add_heading_row(table, text, fill, size):
        tr = table._tbl.add_tr()
        for grid_col in table._tbl.tblGrid.gridCol_lst:
            tr.add_tc()
        cells = row_cells(table, tr)
        merged = cells[0].merge(cells[-1])
        set_text(merged, text, bold=True, size=size)
        shade(merged, fill)

    def add_data_row(table, row):
        tr = table._tbl.add_tr()
        for grid_col in table._tbl.tblGrid.gridCol_lst:
            tr.add_tc()
        values = []
        if flat:
            values += [row.get("_main", ""), row.get("_sub", "")]
        values.append(row.get("product_name", ""))
        values += [format_value(key, row.get(key)) for key in col_keys]
        for i, (cell, value) in enumerate(zip(row_cells(table, tr), values)):
            set_text(cell, value, right=(i in numeric_cols))

    def finish_table(table):
        for i, width in enumerate(col_widths_mm):
            table.columns[i].width = Mm(width)
        for tr in table._tbl.tr_lst:
            for tc, width in zip(tr.tc_lst, col_widths_mm):
                tc.width = Mm(width)

    if flat:
        table = new_table()
        for row in spec.rows:
            add_data_row(table, row)
        finish_table(table)
    else:
        table = None
        current_main = current_sub = None
        for row in spec.rows:
            if row.get("_main") != current_main or table is None:
                if table is not None:
                    finish_table(table)
                    if spec.page_break_per_main:
                        doc.add_page_break()
                    else:
                        doc.add_paragraph()
                current_main, current_sub = row.get("_main"), None
                table = new_table()
                add_heading_row(table, current_main or "", "D9D9D9", 9)
            if row.get("_sub") != current_sub:
                current_sub = row.get("_sub")
                add_heading_row(table, current_sub or "", "F2F2F2", 8)
            add_data_row(table, row)
        if table is not None:
            finish_table(table)

    if spec.summary:
        doc.add_paragraph()
        for line in summary_lines(spec.summary):
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.bold = True
            run.font.size = Pt(9)
            p.paragraph_format.space_after = Pt(0)

    doc.save(path)
